from types import SimpleNamespace

import pytest
from docx import Document as DocxDocument

from backend.app.services.ask import match_quotes, parse_llm_response
from backend.app.services.chunking import Chunk, split_sections, split_text
from backend.app.services.extraction import Section, extract_sections


def test_split_text_normalizes_whitespace():
    assert split_text("Привет,\n\n   мир!\t") == ["Привет, мир!"]


def test_split_text_splits_long_text_into_overlapping_chunks():
    text = " ".join(f"Предложение номер {i}." for i in range(500))

    chunks = split_text(text)

    assert len(chunks) > 1
    assert chunks[0].startswith("Предложение номер 0.")
    assert chunks[-1].endswith("Предложение номер 499.")


def test_split_text_empty():
    assert split_text("   ") == []


def test_split_sections_keeps_page_numbers():
    sections = [Section("Первая страница.", page=1), Section("   ", page=2), Section("Третья.", page=3)]

    assert split_sections(sections) == [
        Chunk("Первая страница.", page=1),
        Chunk("Третья.", page=3),
    ]


def test_get_text_txt(tmp_path):
    path = tmp_path / "file.txt"
    path.write_text("Текст файла", encoding="utf-8")

    assert extract_sections(path) == [Section("Текст файла")]


def test_get_text_docx(tmp_path):
    path = tmp_path / "file.docx"
    document = DocxDocument()
    document.add_paragraph("Первый абзац")
    document.add_paragraph("")
    document.add_paragraph("Второй абзац")
    document.save(path)

    assert extract_sections(path) == [Section("Первый абзац\nВторой абзац")]


def test_get_text_pdf_keeps_page_numbers(tmp_path):
    path = tmp_path / "file.pdf"
    path.write_bytes(make_pdf(["First page", "", "Third page"]))

    assert extract_sections(path) == [
        Section("First page", page=1),
        Section("Third page", page=3),
    ]


def test_get_text_unknown_binary(tmp_path):
    path = tmp_path / "file.bin"
    path.write_bytes(bytes(range(256)) * 4)

    assert extract_sections(path) is None


def make_pdf(pages: list[str]) -> bytes:
    """Минимальный валидный PDF: по строке текста на страницу (пустая — без текста)."""
    # 1 — каталог, 2 — список страниц, 3 — шрифт, дальше пары «страница, контент».
    page_ids = [4 + 2 * index for index in range(len(pages))]
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [%s] /Count %d >>"
        % (b" ".join(b"%d 0 R" % id for id in page_ids), len(pages)),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for page_id, text in zip(page_ids, pages):
        stream = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents %d 0 R /Resources << /Font << /F1 3 0 R >> >> >>" % (page_id + 1)
        )
        objects.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream))

    pdf = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += b"%d 0 obj\n%s\nendobj\n" % (number, body)

    xref = len(pdf)
    pdf += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    pdf += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    pdf += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    return pdf


def chunk(id: int, text: str):
    return SimpleNamespace(id=id, text=text, page=None)


def test_match_quotes_finds_chunks_and_sorts_by_distance():
    chunks = [
        (chunk(1, "Гарантия составляет семь лет."), 0.5),
        (chunk(2, "Директор — Анна Петрова."), 0.2),
    ]

    sources = match_quotes(
        ["семь лет", "Анна Петрова", "выдуманная цитата"],
        chunks,
    )

    assert [source["chunk_id"] for source in sources] == [2, 1]
    assert sources[1]["quote"] == "семь лет"


def test_match_quotes_uses_each_chunk_once_and_at_most_three_quotes():
    chunks = [(chunk(i, f"общий текст {i}"), i / 10) for i in range(5)]

    sources = match_quotes(["общий текст"] * 5, chunks)

    assert [source["chunk_id"] for source in sources] == [0, 1, 2]


def test_parse_llm_response():
    content = '{"answer": "Да", "quotes": ["x"]}'

    assert parse_llm_response(content) == {"answer": "Да", "quotes": ["x"]}


@pytest.mark.parametrize("content", [None, ""])
def test_parse_llm_response_empty(content):
    with pytest.raises(ValueError):
        parse_llm_response(content)
