from dataclasses import dataclass
from pathlib import Path

import magic
import pypdfium2
import pytesseract
from bs4 import BeautifulSoup
from docx import Document
from PIL import Image
from pypdf import PdfReader

DOCX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
IMAGE_MIME_TYPES = {"image/png", "image/jpeg", "image/tiff", "image/bmp", "image/webp"}
OCR_LANGUAGES = "rus+eng"
# Масштаб рендера страницы PDF для OCR: 72 dpi * 3 ≈ 216 dpi.
PDF_RENDER_SCALE = 3


@dataclass
class Section:
    text: str
    page: int | None = None


def ocr(image: Image.Image) -> str:
    return pytesseract.image_to_string(image, lang=OCR_LANGUAGES).strip()


def extract_txt(file_path: Path) -> list[Section]:
    return [Section(file_path.read_text(encoding="utf-8-sig"))]


def extract_html(file_path: Path) -> list[Section]:
    soup = BeautifulSoup(file_path.read_bytes(), "html.parser")

    for tag in soup(["script", "style", "noscript", "template"]):
        tag.decompose()

    return [Section(soup.get_text(separator="\n", strip=True))]


def extract_pdf(file_path: Path) -> list[Section]:
    reader = PdfReader(file_path)
    sections = []
    scanned = None

    for number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()

        if not text:
            # Нет текстового слоя — скорее всего, скан. Распознаём картинку страницы.
            scanned = scanned or pypdfium2.PdfDocument(file_path)
            image = scanned[number - 1].render(scale=PDF_RENDER_SCALE).to_pil()
            text = ocr(image)

        if text:
            sections.append(Section(text, page=number))

    if scanned is not None:
        scanned.close()

    return sections


def extract_image(file_path: Path) -> list[Section]:
    with Image.open(file_path) as image:
        return [Section(ocr(image))]


def extract_docx(file_path: Path) -> list[Section]:
    document = Document(file_path)
    return [
        Section(
            "\n".join(
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text
            )
        )
    ]


def extract_sections(file_path: str | Path) -> list[Section] | None:
    path = Path(file_path)
    mime_type = magic.from_file(str(path), mime=True)

    if mime_type == "text/html":
        return extract_html(path)
    if mime_type.startswith("text/"):
        return extract_txt(path)
    if mime_type == "application/pdf":
        return extract_pdf(path)
    if mime_type in IMAGE_MIME_TYPES:
        return extract_image(path)
    if mime_type == DOCX_MIME_TYPE:
        return extract_docx(path)
    return None
